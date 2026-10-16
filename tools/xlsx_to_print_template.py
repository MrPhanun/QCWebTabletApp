"""Convert "Template Print Spec.xlsx" into spec-print.js, the HTML print layout used by the SPEC page.

The layout keeps the workbook's column widths, row heights, merged cells, fonts and borders, and puts
{{placeholders}} in the cells that SPEC data fills (see DATA). After editing the template, run again:

    python tools/xlsx_to_print_template.py
"""
import html
import json
import os

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'Template Print Spec.xlsx')
OUT = os.path.join(ROOT, 'public', 'js', 'spec-print.js')
LAST_COL, LAST_ROW = 'AV', 43          # print area A1:AV43

# Cell range -> (placeholder, CSS class for the filled value)
DATA = {
    'A7:H7': ('customer', 'v-big'),
    'I7:T7': ('partNo', 'v-big'),
    'S9:AB10': ('header', 'v-big'),     # beside "1) MATERIAL CHECK" (the sample sheet shows the Part No here)
    'C12:AE14': ('wire', 'v-text'),
    'C16:AE18': ('terminal', 'v-text'),
    'C20:AE23': ('housing', 'v-text'),
    'C25:AE26': ('other', 'v-text'),
}
for i in range(1, 12):
    r = 29 + i
    DATA[f'C{r}:G{r}'] = (f'p{i}', 'v-pos')
    DATA[f'H{r}:J{r}'] = (f's{i}', 'v-spec')
    DATA[f'AA{r}:AE{r}'] = (f'p{i + 11}', 'v-pos')
    DATA[f'AF{r}:AH{r}'] = (f's{i + 11}', 'v-spec')

BORDER = {'thin': '1px solid', 'hair': '1px solid', 'medium': '2px solid', 'thick': '3px solid', 'double': '3px double',
          'dotted': '1px dotted', 'dashed': '1px dashed', 'mediumDashed': '2px dashed', 'dashDot': '1px dashed',
          'mediumDashDot': '2px dashed', 'dashDotDot': '1px dashed', 'mediumDashDotDot': '2px dashed', 'slantDashDot': '2px dashed'}
FONTS = {'宋体': "SimSun, 'Times New Roman', serif"}


def main():
    ws = openpyxl.load_workbook(SRC).active
    ncol, nrow = column_index_from_string(LAST_COL), LAST_ROW

    # Column widths (Excel characters -> px) and row heights (pt -> px)
    default_w = ws.sheet_format.defaultColWidth or 8.43
    widths = {c: default_w for c in range(1, ncol + 1)}
    for dim in ws.column_dimensions.values():
        if dim.width:
            for c in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
                if c <= ncol:
                    widths[c] = dim.width
    # Excel's stored width already includes cell padding: px = trunc((256*w + trunc(128/7)) / 256 * 7)
    col_px = [int((256 * widths[c] + 18) / 256 * 7) for c in range(1, ncol + 1)]
    default_h = ws.sheet_format.defaultRowHeight or 15
    row_px = [round((ws.row_dimensions[r].height or default_h) * 4 / 3, 2) for r in range(1, nrow + 1)]

    # Merged ranges, plus the data ranges (which must not clash with template text or merges)
    merges = {}
    for mr in ws.merged_cells.ranges:
        c1, r1, c2, r2 = mr.min_col, mr.min_row, mr.max_col, mr.max_row
        if r1 <= nrow and c1 <= ncol:
            merges[(r1, c1)] = (min(r2, nrow), min(c2, ncol))
    data_at = {}
    for ref, (key, cls) in DATA.items():
        c1, r1, c2, r2 = range_boundaries(ref)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                v = ws.cell(r, c).value
                assert v is None, f'{ref} overlaps template text in {get_column_letter(c)}{r}: {v!r}'
        if (r1, c1) in merges:
            assert merges[(r1, c1)] == (r2, c2), f'{ref} differs from the merged range in the template'
        else:
            for (mr1, mc1), (mr2, mc2) in merges.items():
                assert mr2 < r1 or mr1 > r2 or mc2 < c1 or mc1 > c2, f'{ref} overlaps a merged range in the template'
            merges[(r1, c1)] = (r2, c2)
        data_at[(r1, c1)] = (key, cls)
    covered = set()
    for (r1, c1), (r2, c2) in merges.items():
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1):
                    covered.add((r, c))

    classes, rows_html = {}, []

    def side(cell, name):
        s = getattr(cell.border, name)
        if not s or not s.style:
            return None
        color = s.color.rgb if s.color is not None and isinstance(s.color.rgb, str) else 'FF000000'
        return f'{BORDER.get(s.style, "1px solid")} #{color[-6:]}'

    for r in range(1, nrow + 1):
        tds = []
        for c in range(1, ncol + 1):
            if (r, c) in covered:
                continue
            r2, c2 = merges.get((r, c), (r, c))
            cell = ws.cell(r, c)
            f, a = cell.font, cell.alignment
            css = [f'font-size:{f.sz or 11}pt']
            if f.b:
                css.append('font-weight:700')
            if f.i:
                css.append('font-style:italic')
            if f.color is not None and isinstance(f.color.rgb, str) and f.color.rgb[-6:] != '000000':
                css.append(f'color:#{f.color.rgb[-6:]}')
            if f.name in FONTS:
                css.append(f'font-family:{FONTS[f.name]}')
            h = a.horizontal or ('right' if isinstance(cell.value, (int, float)) else 'left')
            css.append('text-align:' + {'centerContinuous': 'center', 'general': 'left', 'fill': 'left', 'justify': 'left', 'distributed': 'center'}.get(h, h))
            css.append('vertical-align:' + {'center': 'middle', 'top': 'top', 'bottom': 'bottom'}.get(a.vertical, 'bottom'))
            if a.wrap_text or r2 > r:
                css.append('white-space:pre-wrap')
            # Like Excel, text may spill into the next cell only when that cell is empty
            if c2 < ncol and (ws.cell(r, c2 + 1).value is not None or (r, c2 + 1) in data_at):
                css.append('overflow:hidden')
            # Outer borders of the (merged) cell: an edge has a border if any cell along it has one
            edges = (('border-left', 'left', [(rr, c) for rr in range(r, r2 + 1)]),
                     ('border-top', 'top', [(r, cc) for cc in range(c, c2 + 1)]),
                     ('border-right', 'right', [(rr, c2) for rr in range(r, r2 + 1)]),
                     ('border-bottom', 'bottom', [(r2, cc) for cc in range(c, c2 + 1)]))
            for prop, name, cells in edges:
                v = next((x for x in (side(ws.cell(*rc), name) for rc in cells) if x), None)
                if v:
                    css.append(f'{prop}:{v}')
            style = ';'.join(css)
            cls = classes.setdefault(style, f'c{len(classes)}')
            attrs = f' class="{cls}'
            text = ''
            if (r, c) in data_at:
                key, dcls = data_at[(r, c)]
                attrs += f' {dcls}'
                text = '{{' + key + '}}'
            elif cell.value is not None:
                v = cell.value
                text = html.escape(str(int(v) if isinstance(v, float) and v.is_integer() else v))
            attrs += '"'
            if r2 > r:
                attrs += f' rowspan="{r2 - r + 1}"'
            if c2 > c:
                attrs += f' colspan="{c2 - c + 1}"'
            tds.append(f'<td{attrs}>{text}</td>')
        rows_html.append(f'<tr style="height:{row_px[r - 1]}px">{"".join(tds)}</tr>')

    css = '\n'.join(f'.spt .{cls}{{{style}}}' for style, cls in classes.items())
    width = sum(col_px)
    doc = (f'<style>\n{css}\n</style>'
           f'<table class="spt" style="width:{width}px"><colgroup>'
           + ''.join(f'<col style="width:{w}px">' for w in col_px)
           + '</colgroup><tbody>' + ''.join(rows_html) + '</tbody></table>')
    with open(OUT, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('// Generated by tools/xlsx_to_print_template.py from "Template Print Spec.xlsx" - do not edit by hand.\n')
        fh.write('window.SPEC_PRINT_TEMPLATE = ' + json.dumps({'width': width, 'html': doc}, ensure_ascii=False) + ';\n')
    print(f'wrote {OUT}: {width}px wide, {round(sum(row_px))}px tall, {len(classes)} styles, {len(DATA)} data cells')


if __name__ == '__main__':
    main()
